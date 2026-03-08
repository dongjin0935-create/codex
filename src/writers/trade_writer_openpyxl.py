from __future__ import annotations

from pathlib import Path

from src.models import OrderData
from src.utils.file_utils import resolve_unique_path
from src.writers.trade_writer_base import TradeWriterBase


class TradeWriterOpenpyxl(TradeWriterBase):
    def write(self, order: OrderData, mapping: dict, template_path: Path, output_dir: Path, ext: str) -> Path:
        if ext.lower() == "xls":
            raise RuntimeError("openpyxl 백엔드는 xls 저장을 지원하지 않습니다. --trade-ext xlsx 사용 필요")
        try:
            from openpyxl import load_workbook  # type: ignore
        except Exception as exc:
            raise RuntimeError("openpyxl 패키지가 필요합니다.") from exc
        wb = load_workbook(template_path)
        trade_map = mapping["trade"]

        for key, rc in trade_map.get("meta_cells", {}).items():
            sheet, pos = rc.split("!", 1)
            ws = wb[sheet]
            row = int(pos.split("C")[0][1:])
            col = int(pos.split("C")[1])
            if key == "project_no":
                ws.cell(row=row, column=col, value=order.project_no)
            elif key == "order_no":
                ws.cell(row=row, column=col, value=order.order_no)
            elif key == "issue_date":
                ws.cell(row=row, column=col, value=order.issue_date.strftime("%Y-%m-%d"))
            elif key == "due_date":
                ws.cell(row=row, column=col, value=order.due_date.strftime("%Y-%m-%d") if order.due_date else "")

        item_idx = 0
        for spec in trade_map["item_sheets"]:
            ws = wb[spec["sheet_name"]]
            for r in range(spec["start_row"], spec["end_row"] + 1):
                if item_idx >= len(order.items):
                    break
                item = order.items[item_idx]
                cols = spec["item_columns"]
                ws.cell(r, cols["순번"], item.seq)
                ws.cell(r, cols["POR NO"], item.por_no)
                ws.cell(r, cols["자재번호"], item.material_no)
                ws.cell(r, cols["품명 및 규격"], item.description)
                ws.cell(r, cols["단위"], item.unit)
                ws.cell(r, cols["발주수량"], float(item.qty))
                ws.cell(r, cols["납품수량"], float(item.qty))
                ws.cell(r, cols["합격수량"], float(item.qty))
                ws.cell(r, cols["단가"], float(item.unit_price))
                ws.cell(r, cols["금액"], float(item.amount))
                item_idx += 1
        if item_idx < len(order.items):
            raise RuntimeError("을지 시트 용량 부족")

        output_name = f"강남거래명세서_{order.issue_date_compact}_{order.project_no}_{order.order_no}.xlsx"
        out_path = resolve_unique_path(output_dir / output_name)
        wb.save(out_path)
        return out_path
