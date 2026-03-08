from __future__ import annotations

from copy import copy
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from config.constants import FALLBACK_CUSTOMER_CODE, FALLBACK_CUSTOMER_NAME, FALLBACK_ITEM_CODE, VAT_RATE
from src.models import OrderData
from src.utils.excel_utils import copy_row_style
from src.utils.file_utils import resolve_unique_path


def _set_value(row: dict[str, object], key: str, value: object) -> None:
    row[key] = value


def _header_key(name: str, idx: int = 0) -> str:
    return f"{name}#{idx}"


def write_erp_file(order: OrderData, mapping: dict, template_path: Path, output_dir: Path) -> Path:
    erp_map = mapping["erp"]
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:
        raise RuntimeError("openpyxl 패키지가 필요합니다.") from exc

    wb = load_workbook(template_path)
    ws = wb[erp_map["sheet_name"]]
    header_map = erp_map["header_map"]
    start_row = erp_map["first_data_row"]

    # clear existing rows
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

    sample_values = erp_map["sample_row_values"]

    for idx, item in enumerate(order.items, start=1):
        row_num = start_row + idx - 1
        ws.insert_rows(row_num)
        copy_row_style(ws, start_row - 1 if start_row > 1 else start_row, row_num)

        # seed static values from sample row
        for header, cols in header_map.items():
            for n, col in enumerate(cols):
                k = _header_key(header, n)
                ws.cell(row=row_num, column=col).value = sample_values.get(k)
                if ws.cell(row=start_row - 1 if start_row > 1 else row_num, column=col).data_type == "f":
                    ws.cell(row=row_num, column=col)._value = copy(ws.cell(row=start_row - 1 if start_row > 1 else row_num, column=col)._value)

        def put(name: str, value: object, occurrence: int = 0) -> None:
            if name in header_map and len(header_map[name]) > occurrence:
                ws.cell(row=row_num, column=header_map[name][occurrence]).value = value

        put("일자", order.issue_date)
        put("순번", idx)
        put("거래처코드", sample_values.get(_header_key("거래처코드", 0)) or FALLBACK_CUSTOMER_CODE)
        put("거래처명", sample_values.get(_header_key("거래처명", 0)) or FALLBACK_CUSTOMER_NAME)
        put("납기일자", order.due_date or order.issue_date, 0)
        put("프로젝트명", order.project_order_key, 0)
        put("프로젝트명", order.project_order_key, 1)
        put("품목코드", sample_values.get(_header_key("품목코드", 0)) or FALLBACK_ITEM_CODE)
        put("품목명", item.material_no)
        put("규격", item.description)
        put("수량", float(item.qty))
        put("단가", float(item.unit_price))
        put("공급가액", float(item.amount))
        vat = (item.amount * VAT_RATE).quantize(Decimal("1"))
        put("부가세", float(vat))
        put("적요", order.memo_key)
        put("납기일자", order.due_date or order.issue_date, 1)

    output_name = f"ERP업로드_{order.issue_date_compact}_{order.project_no}_{order.order_no}.xlsx"
    out_path = resolve_unique_path(output_dir / output_name)
    wb.save(out_path)
    return out_path
