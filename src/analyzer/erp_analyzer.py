from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from config.constants import ERP_HEADER_SCAN_MAX_ROW, ERP_PRIMARY_SHEET_NAME, ERP_REQUIRED_HEADERS
from src.utils.excel_utils import build_header_map, find_first_data_row, find_header_row


@dataclass(slots=True)
class ErpAnalysis:
    sheet_name: str
    header_row: int
    header_map: dict[str, list[int]]
    first_data_row: int
    sample_row_values: dict[str, object]



def analyze_erp_template(template_path: Path) -> ErpAnalysis:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:
        raise RuntimeError("openpyxl 패키지가 필요합니다.") from exc

    wb = load_workbook(template_path)
    ws = wb[ERP_PRIMARY_SHEET_NAME] if ERP_PRIMARY_SHEET_NAME in wb.sheetnames else None

    if ws is None:
        for candidate in wb.worksheets:
            try:
                header_row = find_header_row(candidate, ERP_REQUIRED_HEADERS, ERP_HEADER_SCAN_MAX_ROW)
                ws = candidate
                break
            except ValueError:
                continue
        if ws is None:
            raise RuntimeError("ERP 샘플에서 입력용 시트를 찾지 못했습니다.")
    else:
        header_row = find_header_row(ws, ERP_REQUIRED_HEADERS, ERP_HEADER_SCAN_MAX_ROW)

    header_map = build_header_map(ws, header_row)
    missing = [h for h in ERP_REQUIRED_HEADERS if h not in header_map]
    if missing:
        raise RuntimeError(f"ERP 샘플에서 필수 헤더를 찾지 못했습니다: {missing}")

    first_data_row = find_first_data_row(ws, header_row)
    sample_row_values: dict[str, object] = {}
    for header, cols in header_map.items():
        for idx, col in enumerate(cols):
            sample_row_values[f"{header}#{idx}"] = ws.cell(row=first_data_row, column=col).value

    return ErpAnalysis(
        sheet_name=ws.title,
        header_row=header_row,
        header_map=header_map,
        first_data_row=first_data_row,
        sample_row_values=sample_row_values,
    )
